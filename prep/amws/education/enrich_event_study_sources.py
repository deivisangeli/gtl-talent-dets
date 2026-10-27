"""Frozen-rule education enrichment; never rewrite estimation inputs or panels."""
import argparse
import collections
import csv
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile

from pilot_io import DATA_ROOT, digest, read_csv, write_csv, write_json

HERE = Path(__file__).resolve().parent
FAMILIES = ('ba_bs', 'ms', 'ma', 'phd')
MS_LABELS = set('ms sm msc scm m8 mjs msce msee msme msche msag msa msctech mspharm mdsc msceng mns msd msph msca masc mst'.split())
EXPECTED_HASH = 'ff51d8001e7bcc764b219eac14cb2081a6dc6f76f998f326491309338cab4bb1'
EARLY_FILES = ('amws_combined_us_geocoded.csv', 'amws_1906_us_unmatched_final.csv',
               'amws_1938_us_unmatched_final.csv', 'amws_1955_us_still_unmatched.csv')
NEW_COLUMNS = [f'{f}_{suffix}' for f in FAMILIES for suffix in
               ('institutions', 'graduation_years', 'graduation_years_raw')]
NEW_COLUMNS += ['education_status', 'education_flags', 'education_parser_version']


def dumps(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


def integer(value):
    if value in (None, '', 'NA', 'NaN'): return None
    try:
        n = float(value)
        return int(n) if n.is_integer() else None
    except (ValueError, OverflowError): return None


def family(award, parser):
    """Use frozen groups plus an explicit master-label whitelist, not new regex."""
    if award['group'] == 'bachelors': return 'ba_bs'
    if award['group'] == 'phd': return 'phd'
    if award['group'] == 'masters_academic':
        label = parser.canon(award['label_as_written'])
        if label in ('ma', 'am'): return 'ma'
        if label in MS_LABELS: return 'ms'
    return None


def project(prediction, parser):
    groups = {f: [] for f in FAMILIES}
    flags = set(prediction['flags'])
    for d in prediction['degrees']:
        f = family(d, parser)
        if f and not d['honorary']:
            groups[f].append(d)
            flags.update(f + ':' + flag for flag in d['flags'])
    fields = {}
    for f, awards in groups.items():
        for suffix, key in [('institutions', 'institution_as_written'),
                            ('graduation_years', 'year'), ('graduation_years_raw', 'year_raw')]:
            fields[f'{f}_{suffix}'] = dumps([d[key] or None for d in awards])
    selected = [d for ds in groups.values() for d in ds]
    fields['education_status'] = ('no_requested_earned_degree_match' if not selected else
                                  'partial_or_flagged' if flags else 'complete')
    fields['education_flags'] = dumps(sorted(flags))
    fields['education_parser_version'] = parser.RULE_VERSION
    return fields, groups


def validate_awards(text, awards, edition, birth_year=None):
    for d in awards:
        assert 0 <= d['start'] <= d['end'] <= len(text), 'Evidence offsets outside biography'
        assert text[d['start']:d['end']] == d['evidence'], 'Evidence substring mismatch'
        assert text[d['label_start']:d['label_end']] == d['label_as_written'], 'Label offsets mismatch'
        for k in ('institution_as_written', 'year_raw'):
            assert not d[k] or d[k] in d['evidence'], f'{k} is outside cited award evidence'
        assert d['year'] is None or 1800 <= d['year'] <= edition, 'Resolved year outside edition'
        if d['year'] is not None and birth_year is not None and d['year'] < birth_year:
            assert 'chronology_conflict' in d['flags'], 'Pre-birth year lacks chronology flag'


def sources(root):
    return [root/'output/amws'/f for f in EARLY_FILES] + [root/'Data/processed/amws/amws_ed86_final.csv']


def source_key(row, file):
    if file.name == 'amws_ed86_final.csv':
        entry_id = integer(row['entry_id'])
        assert entry_id is not None, 'Invalid 1986 entry ID'
        return '1986:' + str(entry_id)
    ed = integer(row['edition']) if file.name == EARLY_FILES[0] else int(file.name.split('_')[1])
    lineid = integer(row['lineid'])
    assert ed in (1906, 1938, 1955) and lineid is not None, 'Invalid edition/source ID'
    return f'{ed}:{lineid}'


def preservation_paths(root, output):
    """Include both old pilots, source files, and AMWS panel outputs recursively."""
    files = set(sources(root))
    for ed in (1906, 1938, 1955):
        files.add(root/'input/amws'/f'amws_{ed}.{ "xlsx" if ed == 1906 else "dta" }')
        files.add(root/'output/amws'/('amws_1955_split_corrected.csv' if ed == 1955 else f'amws_{ed}_cleaned_corrected.csv'))
    for folder in (root/'Data/processed/amws', root/'Data/intermediary/amws', root/'output/amws/education_pilot'):
        files.update(p for p in folder.rglob('*') if p.is_file())
    for p in (root/'output').rglob('*.csv'):
        if output in p.parents or 'education_pilot' in p.parts: continue
        low = p.name.lower()
        if 'amws' in low and ('panel' in low or 'county' in low or 'state' in low): files.add(p)
    files.add(HERE/'education_regex.py')
    return sorted(files)


def load_parser(root):
    pilot = root/'output/amws/education_pilot/multi_edition_regex_pilot'
    manifest = json.loads((pilot/'freeze_manifest.json').read_text(encoding='utf-8'))
    frozen = pilot/'frozen_education_regex.py'
    assert manifest['version'] == 'multi_v03'
    assert digest(frozen) == manifest['parser_sha256'] == EXPECTED_HASH
    assert digest(HERE/'education_regex.py') == EXPECTED_HASH, 'Current parser differs from freeze'
    spec = importlib.util.spec_from_file_location('amws_frozen_education', frozen)
    parser = importlib.util.module_from_spec(spec)
    prior_bytecode = sys.dont_write_bytecode
    try:
        sys.dont_write_bytecode = True
        spec.loader.exec_module(parser)
    finally: sys.dont_write_bytecode = prior_bytecode
    return parser, manifest, frozen


def raw_early(root, stage, rscript):
    dest = stage/'ingestion'
    subprocess.run([str(rscript), str(HERE/'ingest_early.R'), str(root), str(dest)], check=True)
    rows = {r['source_key']: r for ed in (1906, 1938, 1955) for r in read_csv(dest/f'source_{ed}.csv')}
    # readxl normalizes embedded newlines/trims end spaces. Keep exact workbook
    # cells for production evidence; validate that decoding lost no other text.
    from openpyxl import load_workbook
    wb = load_workbook(root/'input/amws/amws_1906.xlsx', read_only=True, data_only=True)
    try:
        for values in wb.worksheets[0].iter_rows(values_only=True):
            lineid = integer(values[0])
            if lineid is None: continue
            r = rows[f'1906:{lineid}']; cell = values[5]
            original = '' if cell is None else str(cell)
            same = re.sub(r'\s+', '', original) == re.sub(r'\s+', '', r['raw_text_adjusted'])
            assert same, f'Workbook decoding changes non-whitespace text at {lineid}'
            r['encoding_note'] = 'exact workbook Unicode cell' + ('; readxl whitespace normalization restored' if original != r['raw_text_adjusted'] else '')
            r['raw_text'] = r['raw_text_adjusted'] = original
    finally: wb.close()
    return rows


def verify_csv_append(baseline, current, excluded_columns=()):
    """Compare original cells and row order; permit only rows appended at the end."""
    with baseline.open(encoding='utf-8-sig', newline='') as old_stream, \
         current.open(encoding='utf-8-sig', newline='') as new_stream:
        old = csv.DictReader(old_stream); new = csv.DictReader(new_stream)
        columns = [c for c in old.fieldnames if c not in excluded_columns]
        assert all(c in new.fieldnames for c in columns), 'Baseline columns missing'
        assert [c for c in new.fieldnames if c in columns] == columns, 'Baseline column order changed'
        baseline_rows = 0
        for row in old:
            candidate = next(new, None)
            assert candidate is not None, 'Baseline rows removed'
            baseline_rows += 1
            assert all(row[c] == candidate[c] for c in columns), \
                f'Baseline row changed at position {baseline_rows}: {current}'
        added_rows = sum(1 for _ in new)
    return dict(baseline_rows=baseline_rows, added_rows=added_rows,
                current_rows=baseline_rows + added_rows)


def verify_pilot_sources(root, split, baseline_dir=None):
    """Allow only a proven append to the 1986 pilot frame; keep its freeze intact."""
    extension = None
    frame = root/'Data/intermediary/amws/amws_ed86_full.csv'
    for filename, checksum in split['source_sha256'].items():
        path = Path(filename)
        current_hash = digest(path)
        if current_hash == checksum:
            continue
        assert baseline_dir is not None and path.resolve() == frame.resolve(), \
            f'Pilot input has changed: {path}'
        baseline = baseline_dir/frame.relative_to(root)
        assert digest(baseline) == checksum, f'Archived pilot baseline differs: {baseline}'
        counts = verify_csv_append(baseline, path)
        assert counts['added_rows'] > 0, 'Changed pilot frame is not an append'
        extension = dict(source=str(path), baseline=str(baseline),
                         baseline_sha256=checksum, current_sha256=current_hash, **counts)
    return extension


def load_previous_release(root, previous_release):
    manifest = json.loads((previous_release/'manifest.json').read_text(encoding='utf-8'))
    assert manifest['parser']['parser_sha256'] == EXPECTED_HASH, 'Previous parser differs'
    assert manifest['parser']['version'] == 'multi_v03'
    expected_names = {p.stem + '_education.csv' for p in sources(root)}
    assert {r['enriched_file'] for r in manifest['sources']} == expected_names
    comparisons = []
    for source in manifest['sources']:
        baseline = previous_release/source['enriched_file']
        assert digest(baseline) == source['output_sha256'], f'Previous enriched file changed: {baseline}'
        current = next(p for p in sources(root) if p.stem + '_education.csv' == baseline.name)
        counts = verify_csv_append(baseline, current, NEW_COLUMNS)
        if current.name != 'amws_ed86_final.csv':
            assert counts['added_rows'] == 0, 'Early-edition source changed'
        comparisons.append(dict(source=str(current), previous_enriched_file=str(baseline), **counts))
    records = read_csv(previous_release/'education_records.csv')
    previous = {r['source_key']: r for r in records}
    assert len(previous) == len(records) == manifest['distinct_source_records'], 'Previous source keys differ'
    return previous, comparisons


def added_coverage(records, new_keys):
    rows = []
    for f in ('all',) + FAMILIES:
        families = FAMILIES if f == 'all' else (f,)
        counts = collections.Counter(records=len(new_keys))
        for key in new_keys:
            row = records[key]
            awards = [award for family_name in families for award in zip(*[
                json.loads(row[f'{family_name}_{suffix}']) for suffix in
                ('institutions', 'graduation_years', 'graduation_years_raw')])]
            counts['identified'] += bool(awards)
            counts['institution'] += any(a[0] for a in awards)
            counts['resolved_year'] += any(a[1] is not None for a in awards)
            counts['complete_tuple'] += any(a[0] and a[1] is not None for a in awards)
            counts['no_requested_earned_degree_match'] += not awards
            counts['multiple_awards'] += len(awards) > 1
            counts['flagged_entries'] += bool(json.loads(row['education_flags']))
        rows.append(dict(edition=1986, family=f, **counts))
    return rows


def run(root, output, rscript, pilot_baseline_dir=None, previous_release=None):
    assert output.parent == root/'output/amws', 'Output must be a distinct AMWS output directory'
    assert not output.exists(), f'Output already exists; preserve it and choose a fresh --output directory: {output}'
    parser, frozen_manifest, frozen = load_parser(root)
    assert (pilot_baseline_dir is None) == (previous_release is None), \
        'Verified refresh requires both --pilot-baseline-dir and --previous-release'
    inputs = sources(root)
    preserve = preservation_paths(root, output)
    if previous_release is not None:
        preserve += sorted(p for p in previous_release.rglob('*') if p.is_file())
        preserve += sorted(p for p in pilot_baseline_dir.rglob('*') if p.is_file())
    before = {str(p): digest(p) for p in preserve}
    # Frozen pilot was evaluated against these raw/canonical sources.
    split = json.loads((frozen.parent/'split_manifest.json').read_text(encoding='utf-8'))
    extension = verify_pilot_sources(root, split, pilot_baseline_dir)
    previous, comparisons = load_previous_release(root, previous_release) if previous_release else ({}, [])
    input_checksums = {str(p): digest(p) for p in inputs}
    output.parent.mkdir(parents=True, exist_ok=True)
    # Keep staging outside the synced Dropbox directory when on the same volume;
    # syncing can transiently lock directories during cleanup or publication.
    temp_root = Path(tempfile.gettempdir())
    staging_parent = temp_root if temp_root.anchor.lower() == output.anchor.lower() else output.parent
    with tempfile.TemporaryDirectory(prefix='amws-education-staging-', dir=staging_parent,
                                     ignore_cleanup_errors=True) as temp:
        stage = Path(temp)
        early = raw_early(root, stage, rscript)
        target_keys = set()
        input_counts = []
        for file in inputs:
            count = 0; keys = set()
            with file.open(encoding='utf-8-sig', newline='') as stream:
                reader = csv.DictReader(stream)
                assert not set(reader.fieldnames) & set(NEW_COLUMNS), 'Education columns already present'
                for row in reader:
                    key = source_key(row, file)
                    if file.name == 'amws_ed86_final.csv': assert key not in keys, 'Duplicate 1986 entry ID'
                    keys.add(key); count += 1
            target_keys.update(keys)
            input_counts.append(dict(source_file=str(file), rows=count, distinct_source_keys=len(keys)))
        missing = sorted(k for k in target_keys if not k.startswith('1986:') and k not in early)
        if missing:
            diagnostic = output.parent/'education_enrichment_join_failures.json'
            write_json(diagnostic, dict(missing_source_keys=missing))
            raise ValueError(f'Unresolved source joins; diagnostics: {diagnostic}')
        assert set(previous) <= target_keys, 'Previous source entries removed'
        new_keys = target_keys - set(previous) if previous_release else set()
        assert all(k.startswith('1986:') for k in new_keys), 'Unexpected new early-edition entries'
        records = {}; counters = collections.defaultdict(collections.Counter)
        award_count = 0
        ledger_fields = ['source_key', 'edition', 'award_index', 'requested_family', 'family_award_index',
                         'group', 'label_as_written', 'institution_as_written', 'year_raw', 'year', 'honorary',
                         'evidence', 'evidence_source', 'start', 'end', 'label_start', 'label_end', 'rule_id', 'flags']
        bio_fields = ['source_key', 'edition', 'birth_year_used', 'source_file', 'source_column',
                      'encoding_note', 'raw_text_adjusted']
        with (stage/'education_awards.csv').open('w', encoding='utf-8-sig', newline='') as ledger_stream, \
             (stage/'source_biographies.csv').open('w', encoding='utf-8-sig', newline='') as bio_stream:
            ledger = csv.DictWriter(ledger_stream, fieldnames=ledger_fields); ledger.writeheader()
            bio = csv.DictWriter(bio_stream, fieldnames=bio_fields); bio.writeheader()

            def parse_record(key, text, birth, source_file, source_column, encoding):
                nonlocal award_count
                ed = int(key.split(':')[0]); pred = parser.extract(text, birth, ed)
                try: validate_awards(text, pred['degrees'], ed, birth)
                except AssertionError as exc: raise ValueError(f'{key}: {exc}') from exc
                fields, groups = project(pred, parser)
                records[key] = dict(source_key=key, edition=ed, birth_year_used=birth,
                                    parser_status=pred['status'], **fields)
                if key in previous:
                    prior = previous[key]
                    assert all(fields[c] == prior[c] for c in NEW_COLUMNS), f'Previous extraction differs: {key}'
                    assert pred['status'] == prior['parser_status']
                    assert birth == integer(prior['birth_year_used']), f'Previous birth year differs: {key}'
                bio.writerow(dict(source_key=key, edition=ed, birth_year_used=birth,
                    source_file=str(source_file), source_column=source_column, encoding_note=encoding, raw_text_adjusted=text))
                family_positions = collections.Counter()
                for index, d in enumerate(pred['degrees'], 1):
                    f = family(d, parser); position = None
                    if f and not d['honorary']: family_positions[f] += 1; position = family_positions[f]
                    ledger.writerow(dict(source_key=key, edition=ed, award_index=index,
                        requested_family=f, family_award_index=position, **{**d, 'flags': dumps(d['flags'])}))
                    award_count += 1
                for f, ds in groups.items():
                    c = counters[(ed, f)]; c['records'] += 1; c['awards'] += len(ds)
                    c['identified'] += bool(ds)
                    c['institution'] += any(d['institution_as_written'] for d in ds)
                    c['resolved_year'] += any(d['year'] is not None for d in ds)
                    c['literal_year'] += any(d['year_raw'] for d in ds)
                    c['complete_tuple'] += any(d['institution_as_written'] and d['year'] is not None for d in ds)
                    c['multiple_awards'] += len(ds) > 1
                    c['partial_or_flagged'] += any(d['flags'] for d in ds)
                    c['unresolved_literal_year'] += any(d['year_raw'] and d['year'] is None for d in ds)
                if len(records) % 10000 == 0: print(f'Extracted {len(records):,} distinct source entries', flush=True)

            for key in sorted(k for k in target_keys if not k.startswith('1986:')):
                r = early[key]; ed = int(r['edition'])
                parse_record(key, r['raw_text_adjusted'], integer(r['birth_year']),
                             root/'input/amws'/f'amws_{ed}.{ "xlsx" if ed == 1906 else "dta" }',
                             {1906: 'workbook_column_6', 1938: 'a', 1955: 'original'}[ed], r['encoding_note'])
            del early
            with inputs[-1].open(encoding='utf-8-sig', newline='') as stream:
                for r in csv.DictReader(stream):
                    parse_record(source_key(r, inputs[-1]), r['raw_text_adjusted'], integer(r['birth_year']),
                                 inputs[-1], 'raw_text_adjusted', 'UTF-8 as read')
        assert set(records) == target_keys
        record_fields = ['source_key', 'edition', 'birth_year_used', 'parser_status'] + NEW_COLUMNS
        write_csv(stage/'education_records.csv', records.values(), record_fields)
        results = []
        for file, counts in zip(inputs, input_counts):
            enriched = stage/(file.stem + '_education.csv')
            with file.open(encoding='utf-8-sig', newline='') as src, enriched.open('w', encoding='utf-8-sig', newline='') as dest:
                reader = csv.DictReader(src); original_fields = reader.fieldnames
                writer = csv.DictWriter(dest, fieldnames=original_fields + NEW_COLUMNS); writer.writeheader()
                for row in reader:
                    row.update({k: records[source_key(row, file)][k] for k in NEW_COLUMNS}); writer.writerow(row)
            count = 0
            with file.open(encoding='utf-8-sig', newline='') as src, enriched.open(encoding='utf-8-sig', newline='') as dest:
                a = csv.DictReader(src); b = csv.DictReader(dest)
                assert b.fieldnames == a.fieldnames + NEW_COLUMNS
                for original in a:
                    enriched_row = next(b)
                    assert original == {k: enriched_row[k] for k in original}, 'Original columns/row order changed'
                    assert all(enriched_row[k] == records[source_key(original, file)][k] for k in NEW_COLUMNS)
                    count += 1
                assert next(b, None) is None and count == counts['rows']
            results.append(dict(**counts, enriched_file=enriched.name, output_sha256=digest(enriched)))
        # Independently reread the emitted ledger and compare every list element.
        reconstructed = collections.defaultdict(lambda: {f: [] for f in FAMILIES})
        with (stage/'education_awards.csv').open(encoding='utf-8-sig', newline='') as stream:
            n = 0
            for row in csv.DictReader(stream):
                n += 1; f = row['requested_family']
                if f and row['honorary'] == 'False':
                    ds = reconstructed[row['source_key']][f]
                    assert int(row['family_award_index']) == len(ds)+1
                    ds.append((row['institution_as_written'] or None, integer(row['year']), row['year_raw'] or None))
            assert n == award_count
        for key, row in records.items():
            for f in FAMILIES:
                arrays = [json.loads(row[f'{f}_{s}']) for s in ('institutions', 'graduation_years', 'graduation_years_raw')]
                assert len({len(a) for a in arrays}) == 1
                assert list(zip(*arrays)) == reconstructed[key][f], f'Ledger/list mismatch at {key}'
        coverage = []
        for (ed, f), c in sorted(counters.items()):
            coverage.append(dict(edition=ed, family=f, **c,
                **{k+'_pct': 100*c[k]/c['records'] for k in ('identified', 'institution', 'resolved_year', 'complete_tuple')}))
        write_csv(stage/'coverage_by_edition.csv', coverage)
        refresh = None
        if previous_release:
            additions = added_coverage(records, new_keys)
            write_csv(stage/'coverage_new_1986.csv', additions)
            refresh = dict(previous_release=str(previous_release), pilot_source_extension=extension,
                           source_comparisons=comparisons, previous_entries_reproduced=len(previous),
                           previous_1986_entries_reproduced=sum(k.startswith('1986:') for k in previous),
                           new_1986_entries=len(new_keys), new_entry_coverage=additions)
            write_json(stage/'refresh_validation.json', refresh)
        after = {str(p): digest(p) for p in preserve}
        assert before == after, 'Protected sources, pilots, parser, or panels changed during run'
        write_json(stage/'preservation_manifest.json', before)
        validation = dict(all_checks_passed=True, frozen_parser_verified=True, complete_source_joins=True,
            original_columns_rows_order_preserved=True, evidence_offsets_verified=True, legal_resolved_years=True,
            ledger_lists_reconciled=True, protected_files_unchanged=True, protected_file_count=len(before),
            distinct_source_records=len(records), award_rows=award_count, enriched_files=results)
        if refresh:
            validation['verified_refresh'] = refresh
        write_json(stage/'validation_report.json', validation)
        runner_snapshot = stage/'enrichment_runner_snapshot.py'
        runner_snapshot.write_bytes(Path(__file__).read_bytes())
        (stage/'frozen_education_regex.py').write_bytes(frozen.read_bytes())
        write_json(stage/'manifest.json', dict(parser=frozen_manifest, frozen_parser_file=str(frozen),
            runner_sha256=digest(runner_snapshot), runner_snapshot_file=runner_snapshot.name,
            master_mapping=dict(ma=['ma','am'], ms=sorted(MS_LABELS)),
            multiple_awards='aligned lists in biography order; no primary award',
            original_sources_unchanged=True, sources=results, input_checksums=input_checksums,
            frozen_parser_snapshot='frozen_education_regex.py', refresh=refresh,
            distinct_source_records=len(records), award_rows=award_count))
        report(stage, results, coverage, validation)
        validate_published(stage)
        # TemporaryDirectory also cleans partial output after any failure. Publish
        # the complete validated directory with one same-filesystem rename.
        stage.rename(output)
    print(dumps(dict(output=str(output), records=len(records), awards=award_count, validation='passed')), flush=True)


def validate_published(output):
    """Independently reread evidence and audit dates; preserve flagged frozen output."""
    bios = {r['source_key']: r for r in read_csv(output/'source_biographies.csv')}
    awards = prebirth = requested_prebirth = 0
    flagged = []
    with (output/'education_awards.csv').open(encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            key = row['source_key']; source = bios[key]
            d = dict(row)
            for f in ('year', 'start', 'end', 'label_start', 'label_end'): d[f] = integer(d[f])
            d['flags'] = json.loads(d['flags'])
            birth = integer(source['birth_year_used'])
            validate_awards(source['raw_text_adjusted'], [d], int(row['edition']), birth)
            awards += 1
            if d['year'] is not None and birth is not None and d['year'] < birth:
                prebirth += 1
                requested_prebirth += bool(row['family_award_index'])
                flagged.append({k: row[k] for k in ('source_key', 'award_index', 'requested_family',
                                                     'label_as_written', 'year', 'honorary', 'flags')})
    result = dict(all_checks_passed=True, ledger_awards_verified=awards,
                  literal_institution_and_year_within_cited_evidence=True,
                  exact_evidence_and_label_offsets=True, prebirth_dates_all_flagged=True,
                  prebirth_ledger_awards=prebirth, prebirth_requested_earned_awards=requested_prebirth,
                  prebirth_cases=flagged)
    write_json(output/'supplemental_validation.json', result)
    report_path = output/'findings_report.txt'
    marker = 'Additional independent evidence/chronology audit:'
    text = report_path.read_text(encoding='utf-8').split(marker)[0].rstrip()
    text += (f'\n\n{marker} all {awards:,} ledger awards have exact offsets and literal institutions/dates within their cited evidence. '
             f'{prebirth} pre-birth dates are flagged chronology_conflict; {requested_prebirth} belongs to a requested earned-degree family. '
             'Explicit four-digit dates remain as reported by the frozen parser despite chronology conflicts; use the flags when selecting dates for analysis.\n')
    report_path.write_text(text, encoding='utf-8')
    return result


def report(stage, results, coverage, validation):
    text = ['AMWS event-study education enrichment — frozen multi_v03', '',
        'Enriched copies preserve original rows, fields, order and selection flags. Existing panels and models are unchanged.',
        f"{validation['distinct_source_records']:,} distinct edition/source entries; {validation['award_rows']:,} ledger rows.",
        'Source rows can repeat geographic candidates; coverage below counts each edition/source ID once.', '',
        'Source copies:']
    text += [f"- {r['enriched_file']}: {r['rows']:,} rows; {r['distinct_source_keys']:,} distinct source IDs" for r in results]
    text += ['', 'Edition | Family | Unique entries | Degree detected | Institution and resolved year | Partial/flagged entries']
    text += [f"{r['edition']} | {r['family']} | {r['records']:,} | {r['identified']:,} ({r['identified_pct']:.1f}%) | {r['complete_tuple']:,} ({r['complete_tuple_pct']:.1f}%) | {r['partial_or_flagged']:,}" for r in coverage]
    if 'verified_refresh' in validation:
        refresh = validation['verified_refresh']
        counts = refresh['new_entry_coverage'][0]
        text += ['', 'Verified updated-final refresh:',
                 f"{refresh['previous_1986_entries_reproduced']:,} prior 1986 entries reproduce their previous extraction exactly; all prior early-edition entries also reproduce exactly.",
                 f"{refresh['new_1986_entries']:,} newly added 1986 entries: {counts['institution']:,} with a requested earned-degree institution, {counts['resolved_year']:,} with a resolved award year, {counts['complete_tuple']:,} with a complete institution/year tuple, {counts['no_requested_earned_degree_match']:,} with no requested earned-degree match.",
                 'New-entry coverage is reported separately in coverage_new_1986.csv; refresh_validation.json records the verified append and baseline comparisons.',
                 'Frozen pilot accuracy figures have not been independently re-estimated for the new OCR entries.']
    text += ['', 'Interpretation:',
        'These are automated extraction coverage counts across the supplied person-source files, including rows excluded upstream. They are not accuracy estimates or counts of people lacking degrees.',
        'All 15 appended columns are described in DATA_DICTIONARY.txt. Partial awards remain in their original list position, with null values for missing fields.',
        'Pilot held-out complete-tuple precision/recall: 1906 94.03%/100%; 1938 98.95%/96.91%; 1955 99.01%/96.15%; 1986 98.96%/95.00%. These combined-family figures do not separately validate MS and MA.',
        'The pilot pooled tuple precision/recall was 98.05%/96.70%. OCR damage, biography boundaries and institution associations remain sources of error. No full-dataset manual verification was performed.',
        'Professional masters, other bachelors, other academic masters, other doctorates and honorary awards remain in the ledger but are excluded from the four requested earned-degree lists.',
        'Reported graduation years mean award years as reported. Ambiguous centuries stay null while literal dates remain available.',
        f"Validation passed, including unchanged checksums of {validation['protected_file_count']} protected source/pilot/panel files."]
    (stage/'findings_report.txt').write_text('\n'.join(text)+'\n', encoding='utf-8')
    dictionary = ['AMWS education columns and joins', '',
        'Enriched copies append the following columns to every original row:',
        '* {ba_bs,ms,ma,phd}_institutions: JSON array of literal institution strings or null.',
        '* {ba_bs,ms,ma,phd}_graduation_years: JSON array of resolved integer award years or null.',
        '* {ba_bs,ms,ma,phd}_graduation_years_raw: JSON array of literal year tokens or null.',
        'Each family has identical array lengths, ordered by biography appearance. Index i in all three arrays describes the same award. [] means no requested earned award matched.',
        '* education_status: no_requested_earned_degree_match, partial_or_flagged, or complete. This is extractor status, not a manual statement about whether the person had degrees.',
        '* education_flags: JSON array of biography flags plus family-prefixed award flags. Use these alongside the detailed ledger when filtering uncertain extractions.',
        '* education_parser_version: multi_v03.', '',
        'Join early copies to education_records.csv by edition:lineid. For unmatched files, edition is the year in the filename. Join 1986 by 1986:entry_id. Preserve duplicate geographic rows and existing kept/inclusion decisions.',
        'education_records.csv has one row per distinct edition/source key. source_biographies.csv preserves decoded text, birth year used for century resolution, source field and decoding decisions.',
        'education_awards.csv contains all parser awards, with award_index per source record and family_award_index for requested earned awards. Array index 0 corresponds to family_award_index 1. Honorary/excluded awards have no family_award_index.',
        'An institution string may name joint institutions exactly as printed; no external institution standardization is applied.',
        'Evidence offsets are zero-based Unicode code-point indices with exclusive end, into source_biographies.csv raw_text_adjusted. raw_text_adjusted is a canonical input alias for early raw biography fields, not rewritten source data.',
        '1986 uses the processed biography and entry_id; full-edition pilot identifiers are a different key and are not interchangeable.',
        'MA includes MA/AM. MS includes the explicit science-master label whitelist in manifest.json, including MSc/ScM/SM and validated OCR variants. MPhil/PhM/MPh and professional masters are ledger-only.',
        'BA/BS and PhD include the validated parser group variants; MB BS medical awards and other doctorates are excluded from these primary lists.',
        'Coverage denominators are unique source entries per edition in the union of enriched source files. They do not apply panel inclusion filters or cross-edition person deduplication.',
        'Use the enriched copies explicitly in new education analyses. Existing event-study read paths remain unchanged.', '',
        'Reproduce from repository root with: python prep/amws/education/enrich_event_study_sources.py --data-root <Dropbox-root> --rscript <Rscript-path>',
        'An existing output directory is never overwritten; --output can specify a new sibling under output/amws.']
    (stage/'DATA_DICTIONARY.txt').write_text('\n'.join(dictionary)+'\n', encoding='utf-8')


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--data-root', type=Path, default=DATA_ROOT)
    ap.add_argument('--output', type=Path)
    ap.add_argument('--rscript', type=Path, default=Path('C:/Program Files/R/R-4.5.2/bin/Rscript.exe'))
    ap.add_argument('--pilot-baseline-dir', type=Path,
                    help='Archived data-root layout proving the 1986 pilot source was only appended')
    ap.add_argument('--previous-release', type=Path,
                    help='Prior enrichment release whose source rows and extractions must reproduce exactly')
    ap.add_argument('--validate-only', action='store_true', help='Audit an existing enriched output without re-extracting')
    args = ap.parse_args()
    root = args.data_root.resolve()
    output = (args.output or root/'output/amws/education_enriched').resolve()
    if args.validate_only: print(dumps(validate_published(output)))
    else: run(root, output, args.rscript,
              args.pilot_baseline_dir.resolve() if args.pilot_baseline_dir else None,
              args.previous_release.resolve() if args.previous_release else None)
